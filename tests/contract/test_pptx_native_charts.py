"""
Real, native PowerPoint charts for statistics slides — ADR-064.

Before this, every "statistics" layout rendered as either plain
centered text, a colored "chip" card (stat_chip=True themes), or an
editorial sidebar panel — never an actual chart object. This covers
the new native-chart path: unit-parsing/chartability logic
(_parse_stat_magnitude, _extract_chartable_stats) and the end-to-end
render (a real GraphicFrame chart shape appears on the slide, with the
right category/value data and original display text preserved as
per-point data labels), plus the explicit non-regressions — mixed
units still fall back to plain text, and chip/editorial themes are
untouched by this entirely.
"""

import io
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from backend.adapters.export.pptx_adapter import (
    PptxExportAdapter, _parse_stat_magnitude, _extract_chartable_stats,
)
from backend.models.recipe import Recipe, Outline, Slide, ContentBlock, BlockType, StructureSource, Theme


def make_recipe(theme_id: str, stat_bullets: list[str]) -> Recipe:
    outline = Outline(structure_source=StructureSource.AI_GENERATED, slides=[
        Slide(order=1, title="Intro", content_blocks=[]),
        Slide(order=2, title="Key Numbers", layout_type="statistics", content_blocks=[
            ContentBlock(type=BlockType.BULLET, text=t) for t in stat_bullets
        ]),
    ])
    return Recipe.new(project_id="p1", source_text="Topic: test", outline=outline,
                       theme=Theme(color_set_id=theme_id), audience_type="general", language="en")


def _chart_shapes(slide):
    return [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.CHART]


# -- _parse_stat_magnitude ---------------------------------------------------

def test_parses_plain_percent():
    assert _parse_stat_magnitude("42%") == ("percent", 42.0)


def test_parses_currency_with_million_suffix():
    unit_kind, value = _parse_stat_magnitude("$320,820M")
    assert unit_kind == "currency"
    assert value == 320820.0 * 1e6


def test_parses_currency_with_billion_suffix_on_the_same_scale_as_million():
    """$1.2B and $320M need to land on a directly comparable numeric
    scale for a bar chart to make sense — this is the whole point of
    folding the suffix into the value instead of treating K/M/B as
    their own unit."""
    billion_kind, billion_value = _parse_stat_magnitude("$1.2B")
    million_kind, million_value = _parse_stat_magnitude("$320M")
    assert billion_kind == million_kind == "currency"
    assert billion_value > million_value
    assert billion_value == 1.2e9


def test_parses_negative_plain_number():
    assert _parse_stat_magnitude("-12") == ("plain", -12.0)


def test_parses_bare_k_suffix_as_plain():
    assert _parse_stat_magnitude("415K") == ("plain", 415_000.0)


def test_returns_none_for_unparseable_text():
    assert _parse_stat_magnitude("") is None
    assert _parse_stat_magnitude("$") is None


# -- _extract_chartable_stats -------------------------------------------------

def test_same_unit_percentages_are_chartable():
    result = _extract_chartable_stats(["42% renewable adoption", "31% wind share", "18% hydro share"])
    assert result is not None
    assert len(result) == 3
    labels, displays, values = zip(*result)
    assert list(displays) == ["42%", "31%", "18%"]
    assert list(values) == [42.0, 31.0, 18.0]
    assert "renewable adoption" in labels[0]


def test_mixed_units_are_not_chartable():
    """A dollar figure and a percentage on the same axis would be
    meaningless — this must opt OUT of charting, not chart something
    misleading."""
    result = _extract_chartable_stats(["$50B market size", "97% satisfaction"])
    assert result is None


def test_single_stat_is_not_chartable():
    """A chart needs something to compare against — one bar alone
    isn't a chart, it's just a number, so this stays below the
    charting threshold regardless of parseability."""
    result = _extract_chartable_stats(["42% renewable adoption"])
    assert result is None


def test_unparseable_stat_falls_back_for_the_whole_slide():
    """One stat that can't be parsed at all (no number the shared
    CHIP_NUMBER_PATTERN can find) takes the whole slide out of
    chart eligibility, rather than silently charting a partial set
    that doesn't match what the person actually wrote."""
    result = _extract_chartable_stats(["42% renewable adoption", "growing fast"])
    assert result is None


# -- End-to-end rendering ------------------------------------------------

def test_same_unit_statistics_render_as_a_real_chart_on_a_plain_theme():
    recipe = make_recipe("neutral", ["42% renewable adoption", "31% wind share", "18% hydro share"])
    output = PptxExportAdapter().export(recipe)
    prs = Presentation(io.BytesIO(output))
    stats_slide = prs.slides[1]
    charts = _chart_shapes(stats_slide)
    assert len(charts) == 1
    chart = charts[0].chart
    plot = chart.plots[0]
    series = plot.series[0]
    assert list(series.values) == [42.0, 31.0, 18.0]
    assert list(chart.plots[0].categories) == ["renewable adoption", "wind share", "hydro share"]
    # Original display text ("42%", not the raw 42.0) must survive as
    # each point's own data label — that's the whole reason for the
    # per-point override instead of a plain number format.
    label_texts = [pt.data_label.text_frame.text for pt in series.points]
    assert label_texts == ["42%", "31%", "18%"]
    # No text-box callouts for this slide — the chart replaces them,
    # doesn't sit alongside them.
    text_boxes = [
        s for s in stats_slide.shapes
        if s.has_text_frame and s.text_frame.text.strip() and s.shape_type != MSO_SHAPE_TYPE.CHART
    ]
    assert not any(tb.text_frame.text.strip() in ("42%", "31%", "18%") for tb in text_boxes)


def test_mixed_unit_statistics_still_fall_back_to_plain_text_unchanged():
    recipe = make_recipe("neutral", ["$50B market size", "97% satisfaction"])
    output = PptxExportAdapter().export(recipe)
    prs = Presentation(io.BytesIO(output))
    stats_slide = prs.slides[1]
    assert len(_chart_shapes(stats_slide)) == 0
    text_boxes = [s for s in stats_slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    all_text = " ".join(tb.text_frame.text for tb in text_boxes)
    assert "$50B market size" in all_text
    assert "97% satisfaction" in all_text


def test_chip_theme_is_completely_unaffected_by_native_charts():
    """stat_chip=True themes (ADR-059) are a deliberate, reference-
    matched visual identity — real charts must never override them,
    even when the underlying stats would otherwise be chartable."""
    recipe = make_recipe("gradient_violet", ["42% renewable adoption", "31% wind share", "18% hydro share"])
    output = PptxExportAdapter().export(recipe)
    prs = Presentation(io.BytesIO(output))
    stats_slide = prs.slides[1]
    assert len(_chart_shapes(stats_slide)) == 0
    from pptx.enum.shapes import MSO_SHAPE
    chips = [
        s for s in stats_slide.shapes
        if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and s.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
    ]
    assert len(chips) == 3


def test_editorial_theme_is_completely_unaffected_by_native_charts():
    """ADR-062's stacked sidebar panel is its own deliberate treatment
    for statistics — not the 'plain, undesigned' case this entry
    upgrades."""
    recipe = make_recipe("editorial_cream", ["42% renewable adoption", "31% wind share", "18% hydro share"])
    output = PptxExportAdapter().export(recipe)
    prs = Presentation(io.BytesIO(output))
    stats_slide = prs.slides[1]
    assert len(_chart_shapes(stats_slide)) == 0
